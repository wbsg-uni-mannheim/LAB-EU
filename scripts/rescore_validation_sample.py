#!/usr/bin/env python3
"""Re-judge the rows of the jurist validation sample and compare with the jurist labels.

Reads one part of 00_Validierung_Stichprobe.xlsx (filled in by the jurist), judges
exactly those (case, arm, criterion) rows again with a given committee and judge
prompt version, and writes a row-level CSV plus a summary. The rubric text is the
judged wording the jurist saw (task snapshot of the August 4 judging), so old and
new verdicts are compared against the same criterion.

Gold label per row: the shown verdict if the jurist marked it correct, the opposite
verdict if marked wrong, none if marked unclear.
"""

from __future__ import annotations

import argparse
import csv
import json
import subprocess
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
PACKET_MAP = REPO_ROOT / "studies" / "de-core-45" / "validation" / "packet_map.tsv"
ARM_RUNS = {
    "Sol · Agent": "agent-gpt-5.6-sol-high/20260730T221733Z",
    "Sol · Baseline": "baseline-gpt-5.6-sol-high/20260730T221734Z",
    "Terra · Agent": "agent-gpt-5.6-terra-high/20260730T221733Z",
    "Terra · Baseline": "baseline-gpt-5.6-terra-high/20260730T221744Z",
    "DeepSeek · Agent": "agent-deepseek-v4-pro-medium/20260730T214602Z",
    "DeepSeek · Baseline": "baseline-deepseek-v4-pro/20260730T221734Z",
}
FLIP = {"pass": "fail", "fail": "pass"}


def load_rows(xlsx: Path, part: int, overrides: dict | None = None) -> list[dict]:
    numbers = {}
    for line in PACKET_MAP.read_text().splitlines():
        _, number, task_id = line.split("\t")
        numbers[number] = task_id.strip()
    ws = load_workbook(xlsx, read_only=True)["Stichprobe"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values))
        if row["Teil"] != part:
            continue
        old = "pass" if row["Urteil"] == "erfüllt" else "fail"
        label = row["Urteil korrekt?"]
        gold = old if label == "ja" else FLIP[old] if label == "nein" else None
        if overrides and str(row["Nr"]) in overrides:
            gold = overrides[str(row["Nr"])]
        rows.append({
            "nr": row["Nr"], "task_id": numbers[row["Prüfpaket"].split(" ", 1)[0]],
            "arm": row["Gezeigte Antwort"], "criterion": row["Kriterium"], "old": old,
            "split": "Gemini" in (row["Stimmen"] or ""), "label": label, "gold": gold,
            "rubric_label": row["Kriterium juristisch korrekt?"],
        })
    return rows


def run_group(args: argparse.Namespace, task_id: str, arm: str, criteria: list[str]) -> Path:
    run_name = task_id.replace("/", "__")
    out = args.output_dir / "scores" / f"{run_name}__{arm.replace(' · ', '-')}.json"
    if out.exists():
        return out
    task_dir = args.tasks_snapshot / task_id
    cmd = [
        args.python, "-m", "evaluation.run", str(task_dir),
        str(args.runs_root / ARM_RUNS[arm] / "tasks" / run_name / "submission"),
        "--rubric", str(task_dir / "evals" / "rubric.json"),
        "--committee-error-retries", "2", "--style-evaluation",
        "--judge-prompt-version", args.prompt_version,
        "--vote-cache-dir", str(args.vote_cache_dir or args.output_dir / "vote-cache"),
        "--output", str(out),
    ]
    if args.judge_model:
        cmd += ["--judge-model", args.judge_model, "--reasoning-effort", args.reasoning_effort,
                "--votes", "1"]
    else:
        cmd += ["--judge-committee", str(args.committee)]
        if not args.full_votes:
            cmd.append("--committee-tiebreaker")
    for cid in criteria:
        cmd += ["--criterion-id", cid]
    out.parent.mkdir(parents=True, exist_ok=True)
    log = out.with_suffix(".log")
    with log.open("w") as fh:
        subprocess.run(cmd, cwd=REPO_ROOT, stdout=fh, stderr=subprocess.STDOUT, check=True)
    return out


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--xlsx", type=Path, required=True)
    p.add_argument("--part", type=int, default=1)
    p.add_argument("--runs-root", type=Path, required=True)
    p.add_argument("--tasks-snapshot", type=Path, required=True,
                   help="Directory that contains de/... task dirs as judged (tasks/ of git archive 1f5be8a).")
    p.add_argument("--committee", type=Path)
    p.add_argument("--judge-model", help="Single judge instead of a committee (OpenAI API).")
    p.add_argument("--reasoning-effort", default="high")
    p.add_argument("--prompt-version", default="v2")
    p.add_argument("--output-dir", type=Path, required=True)
    p.add_argument("--jobs", type=int, default=6)
    p.add_argument("--gold-overrides", type=Path,
                   help="JSON with {\"gold\": {nr: pass|fail|null}} from later jurist clarifications.")
    p.add_argument("--full-votes", action="store_true",
                   help="All three committee members vote on every row (majority of three).")
    p.add_argument("--vote-cache-dir", type=Path,
                   help="Share a vote cache between runs. Defaults to <output-dir>/vote-cache.")
    p.add_argument("--python", default=str(REPO_ROOT / "env" / "bin" / "python"),
                   help="Interpreter with the evaluator dependencies.")
    args = p.parse_args()

    overrides = json.loads(args.gold_overrides.read_text())["gold"] if args.gold_overrides else None
    rows = load_rows(args.xlsx, args.part, overrides)
    groups: dict[tuple[str, str], list[str]] = defaultdict(list)
    for row in rows:
        groups[(row["task_id"], row["arm"])].append(row["criterion"])
    with ThreadPoolExecutor(args.jobs) as pool:
        futures = {key: pool.submit(run_group, args, *key, crits) for key, crits in groups.items()}
        results = {}
        for key, fut in futures.items():
            score = json.loads(fut.result().read_text())
            for r in score["criteria_results"]:
                results[(*key, r["id"])] = r

    out_rows = []
    for row in rows:
        r = results[(row["task_id"], row["arm"], row["criterion"])]
        bases = Counter(
            c.get("basis") for v in r.get("votes", []) for c in v.get("component_checks") or []
        )
        out_rows.append({
            **row, "new": r["verdict"], "new_votes": r.get("vote_counts"),
            "new_tiebreak": r.get("tiebreak_used"), "implied_components": bases.get("implied", 0),
            "new_reasoning": r.get("reasoning", ""),
        })
    with (args.output_dir / "rows.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out_rows[0]))
        w.writeheader()
        w.writerows(out_rows)

    def tally(key: str) -> dict:
        labelled = [r for r in out_rows if r["gold"]]
        return {
            "correct": sum(r[key] == r["gold"] for r in labelled),
            "labelled": len(labelled),
            "too_strict": sum(r[key] == "fail" and r["gold"] == "pass" for r in labelled),
            "too_lenient": sum(r[key] == "pass" and r["gold"] == "fail" for r in labelled),
            "errors": sum(r[key] not in ("pass", "fail") for r in out_rows),
        }

    summary = {
        "old": tally("old"), "new": tally("new"),
        "changed": sum(r["old"] != r["new"] for r in out_rows),
        "unclear_rows": [r["nr"] for r in out_rows if not r["gold"]],
    }
    (args.output_dir / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False))
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
